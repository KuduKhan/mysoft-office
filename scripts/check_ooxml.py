"""Verify generated Office documents with independent format readers."""
from pathlib import Path
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

root=Path(".artifacts/exports")
doc=Document(root/"mysoft.docx")
paragraphs=[p.text for p in doc.paragraphs]
assert "Interop Heading" in paragraphs, paragraphs
assert any("Emphasized" in p for p in paragraphs), paragraphs
assert doc.tables and doc.tables[0].cell(0,0).text=="Cell A", "Document table not readable"
links=[rel.target_ref for rel in doc.part.rels.values() if rel.reltype.endswith("/hyperlink")]
assert "https://example.org/help?a=1&b=two" in links, links
assert any(run.font.superscript for para in doc.paragraphs for run in para.runs), "Superscript run lost"
print("DOCX opens with python-docx, including hyperlinks, superscript and tables")

wb=load_workbook(root/"mysoft.xlsx", data_only=False)
assert wb.sheetnames==["Finance","Overview"],wb.sheetnames
assert wb["Finance"]["B2"].value==345,wb["Finance"]["B2"].value
assert wb["Finance"]["C2"].value=="=B2*2",wb["Finance"]["C2"].value
assert wb["Finance"]["A1"].font.bold is True,"Bold formatting missing"
print("XLSX opens with openpyxl, including sheets, formulas, styles")

ppt=Presentation(root/"mysoft.pptx")
assert len(ppt.slides)==2,len(ppt.slides)
all_text=[[shape.text for shape in slide.shapes if shape.has_text_frame] for slide in ppt.slides]
assert any("Interop Title" in text for text in all_text[0]),all_text
assert any("Follow up" in text for text in all_text[1]),all_text
assert "Presenter briefing" in ppt.slides[0].notes_slide.notes_text_frame.text, "PowerPoint speaker notes lost"
print("PPTX opens with python-pptx, including two slides, editable text and speaker notes")
