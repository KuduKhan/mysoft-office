"""Create real Office-compatible compressed fixtures with independent libraries."""
from pathlib import Path
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(".artifacts/fixtures")
root.mkdir(parents=True, exist_ok=True)

doc = Document()
doc.add_heading("External Document", level=1)
p = doc.add_paragraph("Text from Word ")
p.add_run("bold").bold = True
table = doc.add_table(rows=1, cols=2)
table.cell(0, 0).text = "Budget"
table.cell(0, 1).text = "345"
# Add a real external Word hyperlink relationship for independent import tests.
link_p = doc.add_paragraph("External reference: ")
relation_id = link_p.part.relate_to("https://example.org/guide?revision=2", RT.HYPERLINK, is_external=True)
hyperlink = OxmlElement("w:hyperlink")
hyperlink.set(qn("r:id"), relation_id)
link_run = OxmlElement("w:r")
link_text = OxmlElement("w:t")
link_text.text = "Read guide"
link_run.append(link_text)
hyperlink.append(link_run)
link_p._p.append(hyperlink)
doc.save(root / "external.docx")

wb = Workbook()
ws = wb.active
ws.title = "Finance"
ws["A1"] = "Category"
ws["B2"] = 345
ws["C2"] = "=B2*2"
other = wb.create_sheet("Second")
other["A1"] = "Separate sheet"
wb.save(root / "external.xlsx")

deck = Presentation()
slide = deck.slides.add_slide(deck.slide_layouts[1])
slide.shapes.title.text = "External Slide"
slide.placeholders[1].text = "First point\nSecond point"
slide.notes_slide.notes_text_frame.text = "External presenter briefing\nSecond instruction"
deck.save(root / "external.pptx")

print("Generated independent DOCX, XLSX and PPTX fixtures")
